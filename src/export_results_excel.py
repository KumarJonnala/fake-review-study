"""Collect the final test results into one Excel workbook.

    python3 src/export_results_excel.py

Writes results/results_summary.xlsx with two sheets:

* SVM_XGBoost -- one row per (model, dataset) from results/{svm,xgboost}/*/final_test_results.json
* RAG         -- one row per (dataset, judge) from results/rag/<dataset>.json, where the
                 judges are the retrieval-only baseline plus every configured LLM

Confusion matrices come from src/evaluation.py as [[TN, FP], [FN, TP]] with fake as the
positive class, and are split into four columns here. Both sheets also carry each dataset's
composition (human real / human fake / synthetic row counts), read from its CSV's `origin`.
BERT is left out until it has results for more than d1.
"""
import json
from pathlib import Path

import pandas as pd
import yaml
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
RESULTS = REPO / "results"
OUTPUT = RESULTS / "results_summary.xlsx"
CLASSICAL_MODELS = ["svm", "xgboost"]
DECIMALS = 4


def load_data_config():
    with open(REPO / "src" / "config" / "config.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f)["data"]


def dataset_composition(data_config):
    """Row counts per dataset: human real, human fake, and everything else is synthetic.

    `origin` is human_real, human_fake, or the name of the generating model.
    """
    composition = {}
    for dataset in data_config["datasets"]:
        origin = pd.read_csv(REPO / data_config["dataset_dir"] / f"{dataset}.csv",
                             usecols=["origin"])["origin"]
        n_real = int((origin == "human_real").sum())
        n_fake = int((origin == "human_fake").sum())
        composition[dataset] = {
            "n_human_real": n_real,
            "n_human_fake": n_fake,
            "n_synthetic": len(origin) - n_real - n_fake,
        }
    return composition


def load_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def confusion_columns(metrics):
    (tn, fp), (fn, tp) = metrics["confusion_matrix"]
    return {"TN": tn, "FP": fp, "FN": fn, "TP": tp}


def classical_rows(composition):
    rows = []
    for model in CLASSICAL_MODELS:
        for dataset in composition:
            path = RESULTS / model / dataset / "final_test_results.json"
            if not path.exists():
                print(f"missing, skipped: {path.relative_to(REPO)}")
                continue
            result = load_json(path)
            metrics = result["test_metrics"]
            params = ", ".join(f"{k}={v}" for k, v in result["best_validation_params"].items())
            rows.append({
                "model": model,
                "dataset": dataset,
                **composition[dataset],
                "n_train": result["n_train"],
                "n_validation": result["n_validation"],
                "n_test": result["n_test"],
                "test_fake_fraction": result["test_fake_fraction"],
                "majority_class_accuracy": result["majority_class_accuracy"],
                "accuracy": metrics["accuracy"],
                "precision": metrics["precision"],
                "recall": metrics["recall"],
                "f1": metrics["f1"],
                "roc_auc": metrics.get("roc_auc"),
                "pr_auc": metrics.get("pr_auc"),
                **confusion_columns(metrics),
                "best_validation_params": params,
            })
    return pd.DataFrame(rows)


def rag_row(dataset, composition, result, judge, metrics, by_origin):
    row = {
        "dataset": dataset,
        "judge": judge,
        **composition,
        "train_rows": result["train_rows"],
        "test_rows": result["test_rows"],
        "k": result["k"],
        "embedding_model": result["embedding_model"],
        "accuracy": metrics["accuracy"],
        "precision": metrics["precision"],
        "recall": metrics["recall"],
        "f1": metrics["f1"],
        **confusion_columns(metrics),
        "unknown_predictions": metrics["unknown_predictions"],
        "scored_rows": metrics["scored_rows"],
    }
    # Origins differ per dataset, so these columns are sparse across the sheet.
    for origin, stats in by_origin.items():
        row[f"acc_{origin}"] = stats["accuracy"]
    return row


def rag_rows(composition):
    rows = []
    for dataset in composition:
        path = RESULTS / "rag" / f"{dataset}.json"
        if not path.exists():
            print(f"missing, skipped: {path.relative_to(REPO)}")
            continue
        result = load_json(path)
        rows.append(rag_row(dataset, composition[dataset], result, "retrieval_only",
                            result["retrieval_only_baseline"], result["retrieval_only_by_origin"]))
        for judge, metrics in result["models"].items():
            rows.append(rag_row(dataset, composition[dataset], result, judge, metrics,
                                metrics.get("by_origin", {})))
    return pd.DataFrame(rows)


def format_sheet(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for i, column in enumerate(sheet.iter_cols(), start=1):
        width = max(len(str(c.value)) for c in column if c.value is not None)
        sheet.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)


def main():
    composition = dataset_composition(load_data_config())
    sheets = {
        "SVM_XGBoost": classical_rows(composition).round(DECIMALS),
        "RAG": rag_rows(composition).round(DECIMALS),
    }
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            format_sheet(writer.sheets[name])
            print(f"{name}: {len(frame)} rows")
    print(f"wrote {OUTPUT.relative_to(REPO)}")


if __name__ == "__main__":
    main()
